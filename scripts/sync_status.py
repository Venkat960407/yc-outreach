"""Write send/reply status into the tracker's user columns from data/status.json.

Only moves a row forward (Ready to review -> Sent -> Replied ...); never overwrites a later status
or anything Venkat typed into Response/Notes himself.

  python scripts/sync_status.py path/to/tracker.xlsx
"""
import json, os, sys
from datetime import datetime
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORDER = ["Waiting on build", "Ready to review", "Verified", "Sent", "Followed up", "Replied",
         "Call scheduled", "No reply", "Skip"]


def main(path):
    status = json.load(open(os.path.join(ROOT, "data", "status.json")))
    wb = load_workbook(path); ws = wb["Tracker"]
    h = {c.value: i + 1 for i, c in enumerate(ws[1])}
    changed = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, h["Company"]).value
        s = status.get(name)
        if not s:
            continue
        cur = ws.cell(r, h["Status"]).value or "Ready to review"
        new = s.get("status")
        if new and (cur not in ORDER or ORDER.index(new) > ORDER.index(cur)):
            ws.cell(r, h["Status"], new); changed += 1
        if s.get("date_sent") and not ws.cell(r, h["Date sent"]).value:
            ws.cell(r, h["Date sent"], datetime.fromisoformat(s["date_sent"]))
        if s.get("response") and not ws.cell(r, h["Response"]).value:
            ws.cell(r, h["Response"], s["response"])
    wb.save(path)
    print(f"updated {changed} statuses")


if __name__ == "__main__":
    main(sys.argv[1])
